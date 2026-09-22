#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Submission-grade DOCX report builder + layered verification gate.

Why this module exists
----------------------
`report_generator.py` produces machine-readable output (JSON / SARIF) and
`report_visualizer.py` produces HTML dashboards. Neither produces what SRC /
CNVD / EDUSRC platforms actually accept: a **DOCX submission** with a fixed
section skeleton, Step-style PoC, and real screenshots.

This module closes that last mile, and enforces the discipline that makes a
report survivable in review:

  1. Layered verification gate (hard gates 0-5 + per-type gates) — a finding
     that cannot pass is NOT written into a report. It stays a lead.
  2. Fixed DOCX skeleton (SRC mode) and a generic-product skeleton (0day mode).
  3. Step spec: Burp-style raw request block + one-line result + real
     screenshot + caption. Screenshots are mandatory; a missing one is
     reported, never silently skipped.
  4. Semantic filename `资产 存在 漏洞类型 漏洞.docx`.
  5. De-AI wording scan on the generated text.

Dependency: python-docx (optional). When absent the module degrades to
`--gate-only` behaviour and says so — it never produces a fake DOCX.

CLI
---
  --emit-template                      write assets/report_template.docx
  --findings <path|-> [--gate-only]    run the gate (and optionally build)
  [--unit example] [--mode src|0day|edu] [--outdir reports]
  [--shots <dir>] [--min-severity medium] [--template <path>]
"""

from __future__ import annotations

import argparse
import datetime as _dt
import json
import os
import re
import sys

from utils import load_json

# ---- Severity / type vocabulary --------------------------------------------
SEV_ORDER = {"low": 0, "medium": 1, "high": 2, "critical": 3}
SEV_CN = {"critical": "严重", "high": "高危", "medium": "中危", "low": "低危"}

TYPE_LABELS = {
    "sqli": "SQL注入", "xss": "XSS", "ssrf": "SSRF", "idor": "越权访问",
    "rce": "远程命令执行", "command_injection": "命令注入", "auth_bypass": "认证绕过",
    "priv_esc": "权限提升", "data_exposure": "敏感数据泄露", "info_leak": "信息泄露",
    "file_upload": "文件上传", "path_traversal": "路径穿越", "logic_flaw": "业务逻辑缺陷",
    "race_condition": "条件竞争", "csrf": "CSRF", "xxe": "XXE", "ssti": "SSTI",
    "deserialization": "反序列化", "weak_credential": "弱口令", "misconfig": "安全配置错误",
    "component_exposure": "组件未授权访问", "open_redirect": "开放重定向",
    "graphql_injection": "GraphQL注入", "nosql_injection": "NoSQL注入",
    "jwt_deep_analysis": "JWT缺陷", "mass_assignment": "批量赋值", "cors_misconfig": "CORS配置缺陷",
    "http_smuggling": "HTTP请求走私", "prototype_pollution": "原型污染",
    # ---- v5.8 AI/LLM domain ----
    "prompt_injection": "提示词注入", "agent_tool_abuse": "Agent工具滥用",
    "system_prompt_leak": "系统提示泄露", "jailbreak": "越狱绕过",
    "llm_data_exposure": "大模型数据泄露", "rag_poisoning": "RAG投毒",
    "memory_poisoning": "记忆污染", "excessive_agency": "过度授权",
    "sandbox_escape": "沙箱逃逸", "llm_supply_chain": "工具描述投毒",
    "llm_output_xss": "输出渲染XSS",
    # ---- v5.9 miniprogram domain ----
    "hardcoded_secret": "硬编码密钥", "mp_api_idor": "小程序接口越权",
    "cloud_db_exposure": "云数据库越权", "cloud_function_abuse": "云函数滥用",
    "mp_login_logic": "小程序登录逻辑缺陷", "mp_payment_logic": "小程序支付逻辑缺陷",
    "mp_package_disclosure": "包信息泄露", "mp_render_injection": "小程序渲染注入",
    # ---- v5.10 Android domain ----
    "android_component_exposure": "导出组件未授权", "android_webview_bridge": "WebView桥缺陷",
    "android_provider_exposure": "ContentProvider暴露", "android_intent_redirect": "Intent重定向",
    "android_binder_privilege": "Binder越权", "android_pendingintent_hijack": "PendingIntent劫持",
    "android_deeplink_hijack": "DeepLink劫持",
    # ---- v5.11 PE domain ----
    "memory_corruption": "内存破坏", "format_string": "格式化字符串",
    "dll_hijacking": "DLL劫持", "missing_mitigation": "保护机制缺失",
}

# Tool values that may NOT back a finding: they are inference, not observation.
BANNED_EVIDENCE_TOOLS = {"llm", "static", "source_read", "scanner", "nuclei", "version", "cve"}
OBSERVING_TOOLS = {"httprequest", "runshell", "browser", "curl", "http_request"}

# ---- Layered verification gate ---------------------------------------------
# 硬门 0-5: every finding must pass all of them (缺一不报).
# 按类型命门: the "key" that is different for each vulnerability class.
TYPE_GATES = {
    "data_exposure": (["pii_fields", "data_sample"], "别人的数据 + ≥3 个敏感字段（非公开/非样例）"),
    "idor": (["ab_proof", "cross_account_proof", "b_resource_proof"], "A/B 交叉证明：A 的资源必须用 B 的凭证读到"),
    "ssrf": (["oob_callback", "internal_echo"], "回显内网数据（云元数据/内网 banner）；仅 DNSLog 不算终点"),
    "auth_bypass": (["ato_proof", "arbitrary_user_proof"], "任意用户可绕过/接管，不能只证明自己能登自己号"),
    "rce": (["command_output"], "真实命令执行回显（RCE 无需“别人的数据”）"),
    "command_injection": (["command_output"], "真实命令执行回显"),
    "file_upload": (["execute_proof", "getshell_proof"], "附件的实际效果：getshell 或服务端解析（纯存储不解析不成立）"),
    "sqli": (["database_proof", "db_name"], "最低共识：注出库名 database()（盲注需时间差 + 库名内容）"),
    "logic_flaw": (["state_change_proof"], "操作真实生效且影响他人业务状态"),
    "race_condition": (["state_change_proof", "concurrency_proof"], "并发结果有多次成功响应 / 数据差异证据"),
    # ---- v5.8 AI/LLM domain ----
    # 语言层的“改写”必须配行为差分，系统层的“执行”必须配落地证据。
    "prompt_injection": (["behavior_delta_proof", "stable_repro_proof", "output_diff"],
                         "行为改写稳定复现（同一 payload ≥3 次结果一致）+ 与对照请求的输出差异"),
    "agent_tool_abuse": (["command_output", "internal_echo", "oob_callback", "file_content_proof"],
                         "落地证据：命令回显 / 内网回显 / 真实 OOB 回调 / 读到的文件内容；模型“声称已执行”不算"),
    "system_prompt_leak": (["prompt_fragment", "tool_inventory"],
                           "System Prompt 原文片段或完整工具清单（非公开可查）"),
    "llm_data_exposure": (["pii_fields", "data_sample", "cross_session_proof"],
                          "跨会话/跨用户的他人数据（PII ≥3 字段）"),
    "rag_poisoning": (["cross_account_proof", "retrieval_citation"],
                      "其他账号的查询命中投毒内容 + 检索引用佐证"),
    "memory_poisoning": (["cross_session_proof", "before_after_proof"],
                         "新开会话仍生效的前后对照"),
    "excessive_agency": (["ab_proof", "state_change_proof"],
                         "低权会话完成了本无权完成的动作且真实生效"),
    "sandbox_escape": (["host_file_proof", "host_identity_proof", "oob_callback"],
                       "宿主文件内容 / 宿主身份 / 沙箱到外部的出站连接证据"),
    # ---- v5.9 miniprogram domain ----
    # 小程序接口本质是 HTTP：越权沿用与 Web 相同的 A/B 交叉硬标准，不放低。
    "hardcoded_secret": (["secret_usable_proof", "api_call_proof"],
                         "密钥能实际调通后端并返回真实数据（不是“包里有个疑似 key”）"),
    "mp_api_idor": (["ab_proof", "cross_account_proof", "b_resource_proof"],
                    "A/B 交叉证明：A 的资源必须用 B 的凭证读到（与 Web idor 同标准）"),
    "cloud_db_exposure": (["cross_user_proof", "anon_read_proof"],
                          "匿名或跨用户读到他人集合文档（脱敏样本 + 权限规则截图）"),
    "cloud_function_abuse": (["cross_user_proof", "state_change_proof"],
                             "云函数越权返回他人数据，或越权动作真实生效"),
    "mp_login_logic": (["ato_proof", "arbitrary_user_proof"],
                       "任意用户可登录/接管，或解出他人手机号（脱敏）"),
    "mp_payment_logic": (["payment_proof", "state_change_proof"],
                         "以非应付金额完成支付，或伪造回调把未付订单置为已付"),
    "mp_render_injection": (["execute_proof"],
                            "真机内真实执行（截图/录屏）；仅服务端回显标签不算"),
    # ---- v5.10 Android domain ----
    # 组件类必须给「可复制粘贴的 ADB 命令 + 实际效果」，“导出了但没反应”只是 lead。
    "android_component_exposure": (["adb_repro_cmd", "effect_proof"],
                                   "ADB 可复现命令 + 实际效果；仅 exported=true 不算"),
    "android_webview_bridge": (["execute_proof", "file_content_proof"],
                               "JS 桥被调用的实际效果（读文件/发请求/拿到内部数据）"),
    "android_provider_exposure": (["cross_app_data_proof", "path_traversal_proof"],
                                  "读到非本应用数据，或穿越出 provider 根目录"),
    "android_intent_redirect": (["redirect_proof"],
                                "跳转到应用内部未导出页面且参数可控"),
    "android_binder_privilege": (["cross_privilege_proof", "data_proof"],
                                 "低权调用高权服务并得到数据或实际效果"),
    "android_pendingintent_hijack": (["hijack_proof"],
                                     "第三方可改写 PendingIntent 目标并触发"),
    "android_deeplink_hijack": (["arbitrary_nav_proof", "param_injection_proof"],
                                "跳转到任意内部页面，或参数被当作 URL/路径使用"),
    # ---- v5.11 PE domain ----
    # 静态可疑点只是 lead；崩溃必须证明「可控」，劫持必须证明「被加载」。
    "memory_corruption": (["crash_repro", "control_proof"],
                          "可复现崩溃 + 执行流可控（EIP/RIP 被输入覆盖）；仅崩溃不够"),
    "format_string": (["leak_output", "control_proof"],
                      "栈数据泄漏输出，或格式化写（%n）可控"),
    "dll_hijacking": (["load_proof"],
                      "测试 DLL 被加载并执行（并注明 劫持/搜索顺序/Phantom 哪一种）"),
    "missing_mitigation": (["mitigation_absent_proof"],
                           "以 dumpbin/pestudio 输出证明 ASLR/DEP/CFG/GS 缺失"),
}

BANNED_IMPACT_WORDS = ["理论上", "可能存在", "可能存在风险", "could be", "should be", "疑似"]

DEFER_KEYS = ("falsification", "negative_control", "gate", "verification")


def _gate_inputs(finding: dict) -> dict:
    """Merge finding / gate / verification sub-objects into one flat lookup."""
    flat = dict(finding)
    for key in DEFER_KEYS:
        sub = finding.get(key)
        if isinstance(sub, dict):
            flat.update({k: v for k, v in sub.items() if k not in flat or not flat[k]})
    return flat


def verify_finding(finding: dict) -> dict:
    """Apply the layered verification gate. Pure, never raises."""
    f = _gate_inputs(finding)
    ev = finding.get("evidence") or {}
    hard, quality = [], []

    def hard_gate(name: str, ok: bool, reason: str, hint: str = "") -> None:
        hard.append({"gate": name, "ok": bool(ok), "reason": "" if ok else reason, "hint": hint})

    # 硬门 0 — 先证伪：必须先假设“这是正常功能”并跑否定实验推翻它。
    falsify = f.get("falsification") or f.get("negative_control")
    hard_gate("硬门0 先证伪", bool(falsify), "缺少先证伪记录",
              "记录否定实验：如果这不是漏洞，最可能的解释是什么，你如何排除")

    # 硬门 1 — PoC 可复现：原始请求块 + 响应，且至少重放 2 次。
    req_ok = bool(ev.get("request")) and bool(ev.get("response"))
    tool = str(ev.get("tool") or "").strip().lower()
    tool_ok = tool not in BANNED_EVIDENCE_TOOLS and (tool in OBSERVING_TOOLS or not tool)
    # replay_count lives canonically inside the evidence block; accept it on the
    # finding too (both shapes are in the wild), and fall back to counting
    # replay_results. Coerce to int so a string "3" still passes.
    replays = ev.get("replay_count", finding.get("replay_count"))
    if replays is None:
        replays = len(ev.get("replay_results") or finding.get("replay_results") or [])
    try:
        replays = int(replays)
    except (TypeError, ValueError):
        replays = 0
    hard_gate("硬门1 PoC 可复现", req_ok and tool_ok and replays >= 2,
              "缺原始请求/响应，或证据来源不是可观测请求（%s），或重放次数 <%d（当前 %d）"
              % (tool or "空", 2, replays),
              "贴 Burp 原始请求块（不要 curl），并至少重放 2 次（evidence.replay_count）")

    # 硬门 2 — 危害是链路终局，不是中间信号。
    impact = f.get("impact") or f.get("demonstrated_impact")
    impact_txt = impact if isinstance(impact, str) else json.dumps(impact or {}, ensure_ascii=False)
    weak = any(w in impact_txt for w in BANNED_IMPACT_WORDS)
    hard_gate("硬门2 危害为链路终局", bool(impact_txt.strip()) and not weak and finding.get("status") == "validated",
              "危害未落地（缺 impact 字段，或含“理论上/疑似”措辞，或 status != validated）",
              "写实际打出来的结果：拿到多少条他人数据 / 状态是否真被改 / 命令是否真执行")

    # 硬门 3 — 服务端 / 权限边界确认，不是浏览器 JS 假象。
    hard_gate("硬门3 服务端边界确认", bool(tool) or bool(ev.get("status_code")),
              "未标注证据来源工具", "标注 tool（httpRequest/curl/runShell）以证明是服务端观察")

    # 硬门 4 — 按类型命门（不同类型的"关键钥匙"不同）。
    spec = TYPE_GATES.get(str(finding.get("type")))
    if spec is None:
        hard_gate("硬门4 类型命门", True, "")
    else:
        keys, hint = spec
        gate_label = TYPE_LABELS.get(str(finding.get("type")), str(finding.get("type")))
        hard_gate("硬门4 类型命门（%s）" % gate_label, bool([k for k in keys if f.get(k)]),
                  "未命中类型命门责任字段 %s" % "/".join(keys), hint)

    # 硬门 5 — 链式追问到终局：同根因接口是否穷举，为何到此为止。
    hard_gate("硬门5 链式追问到终局",
              bool(f.get("chain_closed") or f.get("exhausted") or f.get("scope_note")),
              "未记录同根因接口穷举情况 / 为何到此为止",
              "说明同类接口是否全测、纵深追到哪一层、为什么停在这里")

    # 质量项 — 不卡报告，写进报告说明即可。
    quality.append({"item": "跨环境/跨账号复现", "ok": bool(f.get("cross_env")), "reason": "未做二次环境复现"})
    adaptive = finding.get("adaptive_metadata") or {}
    quality.append({"item": "WAF/防护说明", "ok": bool(f.get("waf_note") or adaptive.get("waf_detected") is not None),
                    "reason": "未说明是否有防护及绕过情况"})

    blocked = [g for g in hard if not g["ok"]]
    return {
        "id": finding.get("id", ""),
        "type": finding.get("type", ""),
        "severity": finding.get("severity", ""),
        "passed": not blocked,
        "hard": hard,
        "quality": quality,
        "blocked_by": [g["gate"] for g in blocked],
    }


# ---- De-AI wording scan -----------------------------------------------------
DEAI_PATTERNS = [
    (r"首先|其次|最后，|综上所述|值得注意的是|也就是说|换句话说", "模板化连接词"),
    (r"——(证明了|即|说明|意味着)", "破折号拖尾解释"),
    (r"极具|天然构成|防骗难度|恶性", "形容词渲染"),
    (r"攻击者可能利用该漏洞造成严重影响|存在安全风险，建议", "正确的废话"),
    (r"下图为浏览器内|该截图展示了", "截图元描述"),
]


def deai_scan(text: str) -> list:
    """Return a list of 'pattern → kind' wording violations (advisory, not blocking)."""
    hits = []
    for pattern, kind in DEAI_PATTERNS:
        m = re.search(pattern, text or "")
        if m:
            hits.append({"matched": m.group(0), "kind": kind})
    return hits


# ---- Helpers ----------------------------------------------------------------
def _asset_of(finding: dict) -> str:
    """Derive the asset name used in the filename / report title."""
    for key in ("asset", "domain", "host"):
        if finding.get(key):
            return str(finding[key])
    target = str(finding.get("target") or "")
    m = re.match(r"^[a-zA-Z]+://([^/]+)", target)
    return (m.group(1) if m else target.split("/")[0]) or "unknown-asset"


def _safe_name(name: str) -> str:
    """Strip characters that are illegal in Windows filenames."""
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip().rstrip(".") or "report"


def semantic_filename(finding: dict) -> str:
    label = TYPE_LABELS.get(str(finding.get("type")), str(finding.get("type") or "漏洞"))
    return _safe_name(f"{_asset_of(finding)} 存在{label}漏洞") + ".docx"


def output_dir(outdir: str, mode: str, unit: str | None) -> str:
    if mode == "0day":
        return os.path.join(outdir, "0day")
    if mode == "edu":
        return os.path.join(outdir, "edu报告")
    return os.path.join(outdir, f"{unit}src" if unit else "src")


# ---- DOCX template ----------------------------------------------------------
def _style_font(style, font: str, size: float, bold: bool) -> None:
    """Force font + colour on a style, including the East-Asian font slot."""
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    style.font.name = font
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)          # kill Word's theme blue
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for slot in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rfonts.set(qn(slot), font)


def emit_template(path: str) -> str:
    """Build the plain, review-friendly DOCX template (all-black, no tables)."""
    from docx import Document
    doc = Document()
    _style_font(doc.styles["Normal"], "微软雅黑", 11, False)
    _style_font(doc.styles["Heading 1"], "微软雅黑", 15, True)
    _style_font(doc.styles["Heading 2"], "微软雅黑", 13, True)
    _style_font(doc.styles["Heading 3"], "微软雅黑", 11, True)
    _style_font(doc.styles["List Bullet"], "微软雅黑", 11, False)
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    doc.save(path)
    return path


# ---- DOCX builder -----------------------------------------------------------
def _finding_text(finding: dict) -> str:
    """Flatten everything that will land in the DOCX, for the de-AI scan."""
    parts = [str(v) for k, v in finding.items() if k not in ("evidence", "replay_results")]
    ev = finding.get("evidence") or {}
    parts += [str(ev.get("request", "")), str(ev.get("response", ""))]
    return "\n".join(parts)


def build_docx(finding: dict, template: str, outpath: str, shots_dir: str | None,
               mode: str) -> dict:
    """Render one submission DOCX. Returns {path, steps, missing_shots, warnings}."""
    from docx import Document
    from docx.shared import Inches

    doc = Document(template)
    asset = _asset_of(finding)
    label = TYPE_LABELS.get(str(finding.get("type")), str(finding.get("type") or "漏洞"))
    severity = SEV_CN.get(str(finding.get("severity")), str(finding.get("severity") or ""))
    warnings, missing_shots = [], []

    def h2(text: str) -> None:
        doc.add_heading(text, level=2)

    def para(text: str, style: str | None = None):
        return doc.add_paragraph(str(text), style=style) if style else doc.add_paragraph(str(text))

    def bullets(items) -> None:
        for it in (items if isinstance(items, (list, tuple)) else [items]):
            if str(it).strip():
                doc.add_paragraph(str(it), style="List Bullet")

    def poc_block(text: str) -> None:
        """Raw request block: monospace, no wrapping tricks, never masked."""
        p = doc.add_paragraph()
        run = p.add_run(str(text).replace("\r\n", "\n"))
        run.font.name = "Consolas"
        from docx.oxml.ns import qn
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")

    steps = finding.get("poc_steps")
    if not steps:
        # Fall back to the single-shot evidence pair, flagged so the analyst
        # knows the step narrative is still thin.
        steps = [{
            "title": "触发请求",
            "note": "在目标上重放该请求。",
            "request": (finding.get("evidence") or {}).get("request", ""),
            "result": finding.get("detection_signal", ""),
        }]
        warnings.append("未提供 poc_steps，已用 evidence 请求块生成单步 PoC（建议补成分步叙事）")

    if mode == "0day":
        _build_generic(doc, finding, asset, label, severity, steps, shots_dir, missing_shots, h2, para, bullets, poc_block, Inches)
    else:
        _build_src(doc, finding, asset, label, severity, steps, shots_dir, missing_shots, h2, para, bullets, poc_block, Inches)

    os.makedirs(os.path.dirname(outpath) or ".", exist_ok=True)
    doc.save(outpath)

    if missing_shots:
        warnings.append("缺真实截图：%s（平台大概率驳回，请补图或说明为何无法用浏览器截图）" % "、".join(missing_shots))
    for hit in deai_scan(_finding_text(finding)):
        warnings.append("去AI腔：命中「%s」→ %s" % (hit["matched"], hit["kind"]))
    return {"path": outpath, "steps": len(steps), "missing_shots": missing_shots, "warnings": warnings}


def _add_steps(doc, steps, shots_dir, missing, h2, para, poc_block, Inches) -> None:
    """Step spec: Heading 3 → note → PoC block → one-line result → screenshot."""
    h2("漏洞复现步骤（POC）")
    for i, step in enumerate(steps, 1):
        doc.add_heading(f"Step {i}：{step.get('title', '')}", level=3)
        if step.get("note"):
            para(step["note"])
        if step.get("request"):
            para("PoC:")
            poc_block(step["request"])
        if step.get("result"):
            para("结果: " + str(step["result"]))

        shot = step.get("shot") or _find_shot(shots_dir, i)
        if shot and os.path.exists(shot):
            para("截图(%s):" % (step.get("shot_note") or step.get("title", "")))
            doc.add_picture(shot, width=Inches(6))
            para("图：" + os.path.basename(shot))
        else:
            missing.append("Step %d" % i)
            para("截图(%s): 【缺真实截图 — 请补 %s/step%d_*.png】"
                 % (step.get("shot_note") or step.get("title", ""), shots_dir or "shots", i))


def _find_shot(shots_dir: str | None, idx: int) -> str | None:
    """Match `step<N>_*.png|jpg` inside the shots directory."""
    if not shots_dir or not os.path.isdir(shots_dir):
        return None
    prefix = f"step{idx}_"
    for fn in sorted(os.listdir(shots_dir)):
        low = fn.lower()
        if low.startswith(prefix) and low.endswith((".png", ".jpg", ".jpeg")):
            return os.path.join(shots_dir, fn)
    return None


def _build_src(doc, f, asset, label, severity, steps, shots_dir, missing, h2, para, bullets, poc_block, Inches) -> None:
    """SRC skeleton — the fixed Heading-2 order, nothing else."""
    h2("漏洞名称")
    para(f"{asset} 存在{label}漏洞，可{str(f.get('impact_short') or '造成实际危害')}")

    h2("漏洞等级")
    para(f"{severity}。{f.get('severity_rationale') or f.get('detection_signal', '')}")

    h2("漏洞类型")
    para(f"{label}（{f.get('type', '')}{('，' + str(f['cwe'])) if f.get('cwe') else ''}）")

    h2("漏洞影响资产")
    para(asset)
    bullets(f.get("related_assets") or [])

    h2("漏洞URL")
    for u in (f.get("urls") or [f.get("target", "")]):
        if str(u).strip():
            para(str(u))

    h2("漏洞描述")
    para(f.get("description") or f.get("detection_signal", ""))
    bullets(f.get("attacker_capability") or [])

    if f.get("session_context"):
        h2("测试账号与会话上下文")
        para(f["session_context"])

    _add_steps(doc, steps, shots_dir, missing, h2, para, poc_block, Inches)

    h2("漏洞危害")
    bullets(f.get("impact") if isinstance(f.get("impact"), (list, tuple)) else [f.get("impact", "")])
    if f.get("data_sample"):
        para("影响数据样本：")
        poc_block(f["data_sample"])
    para("边界声明：未执行超出该漏洞验证所需的任何写操作/越权动作。")

    h2("修复建议")
    bullets(f.get("remediation") if isinstance(f.get("remediation"), (list, tuple))
            else [f.get("remediation", "见漏洞描述根因，按最小权限与输入校验原则修复")])


def _build_generic(doc, f, asset, label, severity, steps, shots_dir, missing, h2, para, bullets, poc_block, Inches) -> None:
    """0day / generic-product skeleton (different section order, same Step spec)."""
    h2("漏洞报告标题")
    para(f"{f.get('vendor', '')} {f.get('product', asset)} 存在{label}漏洞".strip())

    h2("漏洞发现时间")
    para(f.get("discovered_at") or _dt.date.today().isoformat())

    h2("漏洞技术类型")
    para(f"{label}（{f.get('type', '')}{('，' + str(f['cwe'])) if f.get('cwe') else ''}）")

    h2("漏洞描述")
    para(f.get("product_intro", ""))
    para(f.get("description") or f.get("detection_signal", ""))

    h2("漏洞危害")
    bullets(f.get("impact") if isinstance(f.get("impact"), (list, tuple)) else [f.get("impact", "")])

    h2("漏洞厂商全称")
    para(f.get("vendor", ""))

    h2("已知受影响产品及版本")
    bullets(f.get("affected_versions") or [])

    h2("互联网资产证明")
    para("测绘语法：" + str(f.get("fofa_query") or f.get("shodan_query") or ""))
    para("独立 IP 数量：" + str(f.get("asset_count") or ""))
    para("统计来源与时间：" + str(f.get("asset_stats_source") or ""))

    h2("1、漏洞技术细节")
    _add_steps(doc, steps, shots_dir, missing, h2, para, poc_block, Inches)
    para("触发条件：" + str(f.get("trigger_condition") or ""))

    h2("2、复现证明")
    bullets(f.get("case_proofs") or [])

    h2("3、修复方案")
    para("厂商修复：" + str(f.get("vendor_fix") or ""))
    para("运维临时方案：" + str(f.get("ops_mitigation") or ""))

    h2("4、备注")
    para(f.get("notes") or "未执行超出验证所需的任何写操作；案例 IP 提交时现场验证存活。")


# ---- Duplicate pre-check ----------------------------------------------------
def dup_warnings(outdir_path: str, finding: dict) -> list:
    """Cheap 4-way duplicate hint: filename asset+type overlap in the same dir."""
    if not os.path.isdir(outdir_path):
        return []
    asset = _asset_of(finding).lower()
    label = TYPE_LABELS.get(str(finding.get("type")), str(finding.get("type") or ""))
    hits = [fn for fn in os.listdir(outdir_path)
            if fn.lower().endswith(".docx") and asset in fn.lower() and label in fn]
    if hits:
        return ["同目录已有同资产同类型报告 %s —— 先做四项查重（资产/根因/接口/影响面），"
                "重合则改为在原报告上补强" % "、".join(hits)]
    return []


# ---- CLI --------------------------------------------------------------------
def _force_utf8_stdout() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass


def _default_asset(name: str) -> str:
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", name)


def main(argv: list | None = None) -> int:
    _force_utf8_stdout()
    ap = argparse.ArgumentParser(description="GKN-Phantom submission DOCX builder")
    ap.add_argument("--emit-template", action="store_true", help="write the DOCX template and exit")
    ap.add_argument("--findings", help="findings JSON file, or '-' for stdin")
    ap.add_argument("--gate-only", action="store_true", help="run the verification gate only")
    ap.add_argument("--template", default=_default_asset("report_template.docx"))
    ap.add_argument("--outdir", default="reports")
    ap.add_argument("--unit", default=None, help="target unit name (reports/<unit>src/)")
    ap.add_argument("--mode", choices=["src", "0day", "edu"], default="src")
    ap.add_argument("--shots", default=None, help="screenshot source directory")
    ap.add_argument("--min-severity", choices=list(SEV_ORDER), default="medium")
    args = ap.parse_args(argv)

    if args.emit_template:
        path = os.path.abspath(args.template) if args.template else _default_asset("report_template.docx")
        try:
            print(json.dumps({"status": "template_written", "path": emit_template(path)}, ensure_ascii=False))
        except ImportError:
            print(json.dumps({"status": "error", "error": "python-docx not installed; "
                              "pip install python-docx"}, ensure_ascii=False))
            return 3
        return 0

    if not args.findings:
        print("error: --findings is required (or use --emit-template)", file=sys.stderr)
        return 2

    data = load_json(args.findings)
    findings = data.get("findings", []) if isinstance(data, dict) else data
    if not isinstance(findings, list):
        print("error: findings payload must be a list or {'findings': [...]}", file=sys.stderr)
        return 2

    floor = SEV_ORDER[args.min_severity]
    gate_results, skipped, builds = [], [], []
    for f in findings:
        sev = str(f.get("severity") or "low")
        if SEV_ORDER.get(sev, 0) < floor:
            skipped.append({"id": f.get("id", ""), "severity": sev, "reason": "低于 --min-severity"})
            continue
        gate_results.append(verify_finding(f))

    passed = [g for g in gate_results if g["passed"]]
    report = {
        "generated_at": _dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "mode": args.mode,
        "min_severity": args.min_severity,
        "total": len(findings),
        "gated": len(gate_results),
        "passed": len(passed),
        "blocked": len(gate_results) - len(passed),
        "skipped_below_floor": skipped,
        "gates": gate_results,
    }

    if args.gate_only:
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0 if passed else 1

    if not passed:
        print(json.dumps({**report, "status": "no_report_written",
                          "reason": "没有任何 finding 通过分层验证门；补齐阻塞项后重跑"},
                         ensure_ascii=False, indent=2))
        return 1

    try:
        import docx  # noqa: F401
    except ImportError:
        print(json.dumps({**report, "status": "error",
                          "error": "python-docx not installed; only --gate-only is available"},
                         ensure_ascii=False, indent=2))
        return 3

    template = os.path.abspath(args.template)
    if not os.path.exists(template):
        emit_template(template)

    outdir = output_dir(args.outdir, args.mode, args.unit)
    by_id = {str(f.get("id", i)): f for i, f in enumerate(findings)}
    for g in sorted(passed, key=lambda x: -SEV_ORDER.get(x["severity"], 0)):
        finding = by_id.get(g["id"], {})
        # Duplicate hint must run BEFORE the new file lands in the directory,
        # otherwise it always matches itself.
        dups = dup_warnings(outdir, finding)
        outpath = os.path.join(outdir, semantic_filename(finding))
        res = build_docx(finding, template, outpath, args.shots, args.mode)
        res["warnings"] = dups + res["warnings"]
        res["id"], res["severity"] = g["id"], g["severity"]
        builds.append(res)

    report["status"] = "reports_written"
    report["output_dir"] = os.path.abspath(outdir)
    report["artifacts"] = builds
    gate_path = os.path.join(outdir, "_gate_report.json")
    os.makedirs(outdir, exist_ok=True)
    with open(gate_path, "w", encoding="utf-8") as fh:
        json.dump(report, fh, ensure_ascii=False, indent=2)
    report["gate_report"] = gate_path
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
